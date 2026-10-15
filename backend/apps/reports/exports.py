"""Render a report payload as an Excel workbook or a PDF."""
import io
from decimal import Decimal

from django.template.loader import render_to_string
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from xhtml2pdf import pisa


def _display(value, column_type):
    if column_type == "money":
        return f"{Decimal(value):,.2f}"
    return value


def to_xlsx(report):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = report["title"][:31]
    columns = report["columns"]

    sheet.append([f"Deep Water Refilling Station - {report['title']}"])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append([f"Period: {report['date_from']} to {report['date_to']}"])
    sheet.append([])
    sheet.append([c["label"] for c in columns])
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0369A1")
        cell.alignment = Alignment(horizontal="center")

    def cell_value(column, value):
        if column["type"] == "money":
            return Decimal(value)
        return value

    for row in report["rows"]:
        sheet.append([cell_value(c, row[c["key"]]) for c in columns])
    if report["totals"]:
        total_row = ["Total"] + [
            cell_value(c, report["totals"][c["key"]]) if c["key"] in report["totals"] else None
            for c in columns[1:]
        ]
        sheet.append(total_row)
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)

    for index, column in enumerate(columns, start=1):
        letter = get_column_letter(index)
        for cell in sheet[letter][header_row:]:
            if column["type"] == "money":
                cell.number_format = "#,##0.00"
        widest = max(len(str(c.value)) if c.value is not None else 0 for c in sheet[letter][header_row - 1 :])
        sheet.column_dimensions[letter].width = min(max(widest + 4, 12), 40)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def to_pdf(report):
    columns = report["columns"]

    def cell(column, value):
        return {"text": _display(value, column["type"]), "right": column["type"] in ("money", "number")}

    rows = [[cell(c, row[c["key"]]) for c in columns] for row in report["rows"]]
    totals = None
    if report["totals"]:
        totals = [{"text": "Total", "right": False}] + [
            cell(c, report["totals"][c["key"]]) if c["key"] in report["totals"] else {"text": "", "right": False}
            for c in columns[1:]
        ]
    html = render_to_string(
        "reports/report.html",
        {
            "report": report,
            "headers": [c["label"] for c in columns],
            "rows": rows,
            "totals": totals,
        },
    )
    buffer = io.BytesIO()
    result = pisa.CreatePDF(html, dest=buffer)
    if result.err:
        raise RuntimeError("Could not render the PDF report.")
    return buffer.getvalue()
