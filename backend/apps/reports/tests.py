import datetime
import io
from decimal import Decimal

from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.core.testing import APITestBase
from apps.orders.models import Order, OrderDetail
from apps.payments.models import Payment

D = datetime.date


class ReportBase(APITestBase):
    def make_order(self, status, delivery_date, lines, rider=None, customer=None):
        """Create an order straight through the ORM (history dates are not allowed via the API)."""
        order = Order.objects.create(
            customer=customer or self.customer,
            rider=rider,
            delivery_date=delivery_date,
            delivery_address="Somewhere",
            order_status=status,
        )
        total = Decimal("0.00")
        for service, quantity in lines:
            subtotal = service.price * quantity
            OrderDetail.objects.create(
                order=order, service=service, quantity=quantity, unit_price=service.price, subtotal=subtotal
            )
            total += subtotal
        order.total_amount = total
        order.save()
        Payment.objects.create(order=order, payment_method="cod", amount=total)
        return order

    def setUp(self):
        r, s = self.round_refill, self.slim_refill  # 25.00 and 30.00
        self.make_order("delivered", D(2026, 9, 10), [(r, 2)], rider=self.rider)  # 50
        self.make_order("delivered", D(2026, 9, 10), [(r, 1), (s, 1)], rider=self.rider)  # 55
        self.make_order("delivered", D(2026, 9, 12), [(s, 3)], rider=self.rider2)  # 90
        self.make_order("delivered", D(2026, 10, 2), [(r, 4)], rider=self.rider)  # 100
        self.make_order("failed", D(2026, 9, 11), [(r, 1)], rider=self.rider2)
        self.make_order("cancelled", D(2026, 9, 10), [(r, 9)])
        self.make_order("pending", D(2026, 9, 13), [(s, 1)])
        self.admin_client = self.as_user(self.admin)

    def get(self, url, **params):
        return self.admin_client.get(url, params)


class ReportPermissionTests(ReportBase):
    urls = [
        "/api/reports/sales/daily/",
        "/api/reports/sales/monthly/",
        "/api/reports/orders-by-status/",
        "/api/reports/deliveries-per-rider/",
        "/api/reports/top-services/",
    ]

    def test_admin_only(self):
        for url in self.urls:
            self.assertEqual(APIClient().get(url).status_code, 401, url)
            for user in (self.staff, self.customer, self.rider_user):
                self.assertEqual(self.as_user(user).get(url).status_code, 403, (url, user.role))
            self.assertEqual(self.get(url).status_code, 200, url)


class SalesReportTests(ReportBase):
    def test_daily_sales_counts_delivered_only_and_fills_gaps(self):
        response = self.get("/api/reports/sales/daily/", date_from="2026-09-09", date_to="2026-09-13")
        self.assertEqual(response.status_code, 200)
        data = response.data
        self.assertEqual(data["report"], "sales_daily")
        self.assertEqual([c["key"] for c in data["columns"]], ["period", "orders", "total_sales"])
        self.assertEqual(
            [(r["period"], r["orders"], r["total_sales"]) for r in data["rows"]],
            [
                ("2026-09-09", 0, "0.00"),
                ("2026-09-10", 2, "105.00"),
                ("2026-09-11", 0, "0.00"),
                ("2026-09-12", 1, "90.00"),
                ("2026-09-13", 0, "0.00"),
            ],
        )
        self.assertEqual(data["totals"], {"orders": 3, "total_sales": "195.00"})

    def test_monthly_sales(self):
        response = self.get("/api/reports/sales/monthly/", date_from="2026-08-01", date_to="2026-10-31")
        self.assertEqual(
            [(r["period"], r["orders"], r["total_sales"]) for r in response.data["rows"]],
            [("2026-08", 0, "0.00"), ("2026-09", 3, "195.00"), ("2026-10", 1, "100.00")],
        )
        self.assertEqual(response.data["totals"]["total_sales"], "295.00")

    def test_default_range_is_the_last_30_days(self):
        data = self.get("/api/reports/sales/daily/").data
        today = timezone.localdate()
        self.assertEqual(data["date_to"], today.isoformat())
        self.assertEqual(data["date_from"], (today - datetime.timedelta(days=29)).isoformat())
        self.assertEqual(len(data["rows"]), 30)

    def test_bad_ranges(self):
        url = "/api/reports/sales/daily/"
        bad_date = self.get(url, date_from="yesterday")
        self.assertEqual(bad_date.status_code, 400)
        self.assertIn("date_from", bad_date.data["errors"])
        self.assertEqual(self.get(url, date_from="2026-10-02", date_to="2026-10-01").status_code, 400)
        self.assertEqual(self.get(url, date_from="2020-01-01", date_to="2026-01-01").status_code, 400)
        self.assertEqual(self.get(url, **{"format": "docx"}).status_code, 400)


class OtherReportTests(ReportBase):
    def test_orders_by_status(self):
        data = self.get("/api/reports/orders-by-status/", date_from="2026-09-01", date_to="2026-09-30").data
        by_status = {r["status"]: r["orders"] for r in data["rows"]}
        self.assertEqual(by_status["delivered"], 3)
        self.assertEqual(by_status["failed"], 1)
        self.assertEqual(by_status["cancelled"], 1)
        self.assertEqual(by_status["pending"], 1)
        self.assertEqual(by_status["scheduled"], 0)
        self.assertEqual(len(data["rows"]), len(Order.Status.choices))
        self.assertEqual(data["totals"]["orders"], 6)

    def test_deliveries_per_rider(self):
        data = self.get("/api/reports/deliveries-per-rider/", date_from="2026-09-01", date_to="2026-09-30").data
        rows = {r["rider_id"]: r for r in data["rows"]}
        first, second = rows[self.rider.pk], rows[self.rider2.pk]
        self.assertEqual((first["assigned"], first["delivered"], first["failed"]), (2, 2, 0))
        self.assertEqual(first["delivered_value"], "105.00")
        self.assertEqual((second["assigned"], second["delivered"], second["failed"]), (2, 1, 1))
        self.assertEqual(data["totals"]["delivered"], 3)

    def test_top_services_ranked_by_quantity(self):
        data = self.get("/api/reports/top-services/", date_from="2026-09-01", date_to="2026-10-31").data
        # delivered: round 2+1+4 = 7 (175.00), slim 1+3 = 4 (120.00)
        self.assertEqual(
            [(r["rank"], r["service"], r["quantity"], r["revenue"]) for r in data["rows"]],
            [(1, "Round 5-gal Refill", 7, "175.00"), (2, "Slim 5-gal Refill", 4, "120.00")],
        )
        self.assertEqual(len(self.get("/api/reports/top-services/", limit="1", date_from="2026-09-01", date_to="2026-10-31").data["rows"]), 1)
        self.assertEqual(self.get("/api/reports/top-services/", limit="0").status_code, 400)


class ReportExportTests(ReportBase):
    params = {"date_from": "2026-09-09", "date_to": "2026-09-13"}

    def test_xlsx_export(self):
        response = self.get("/api/reports/sales/daily/", format="xlsx", **self.params)
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn("sales_daily_2026-09-09_to_2026-09-13.xlsx", response["Content-Disposition"])
        sheet = load_workbook(io.BytesIO(response.content)).active
        values = [[c.value for c in row] for row in sheet.iter_rows()]
        self.assertIn(["Date", "Orders", "Total sales (PHP)"], values)
        day = next(v for v in values if v[0] == "2026-09-10")
        self.assertEqual((day[1], Decimal(str(day[2]))), (2, Decimal("105.00")))
        self.assertEqual(values[-1][0], "Total")
        self.assertEqual(Decimal(str(values[-1][2])), Decimal("195.00"))

    def test_pdf_export(self):
        response = self.get("/api/reports/sales/daily/", format="pdf", **self.params)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))
        self.assertIn(".pdf", response["Content-Disposition"])

    def test_every_report_exports_both_formats(self):
        for url in ReportPermissionTests.urls:
            for fmt, magic in (("pdf", b"%PDF"), ("xlsx", b"PK")):
                response = self.get(url, format=fmt, **self.params)
                self.assertEqual(response.status_code, 200, (url, fmt))
                self.assertTrue(response.content.startswith(magic), (url, fmt))

    def test_empty_report_still_exports(self):
        response = self.get("/api/reports/top-services/", format="pdf", date_from="2000-01-01", date_to="2000-01-31")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b"%PDF"))
