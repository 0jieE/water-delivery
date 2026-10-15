import io
from decimal import Decimal

from django.core import mail
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.core.testing import APITestBase, tomorrow
from apps.deliveries.models import CashRemittance
from apps.notifications.models import Notification
from apps.orders.models import Order
from apps.payments.models import Payment

HELD = "/api/deliveries/cash/"
REMIT = "/api/deliveries/cash/remit/"
HISTORY = "/api/deliveries/cash/remittances/"
MINE = "/api/deliveries/my/cash/"


class CashBase(APITestBase):
    def deliver(self, rider_user, rider, quantity=None, method="cod", collected=True, finish="delivered"):
        """Take an order through the whole flow with the API; returns the order id."""
        items = [{"service": self.round_refill.pk, "quantity": quantity}] if quantity else None
        order = self.place_order(**({"items": items} if items else {}), payment_method=method)
        url = f"/api/orders/{order['id']}"
        staff = self.as_user(self.staff)
        staff.post(f"{url}/confirm/")
        staff.post(
            f"{url}/schedule/", {"delivery_date": tomorrow().isoformat(), "rider": rider.pk}, format="json"
        )
        client = self.as_user(rider_user)
        client.post(f"/api/deliveries/my/{order['id']}/status/", {"status": "out_for_delivery"}, format="json")
        if finish:
            body = {"status": finish, "cod_collected": collected, "remarks": "x"}
            response = client.post(f"/api/deliveries/my/{order['id']}/status/", body, format="json")
            assert response.status_code == 200, response.data
        return order["id"]

    def setUp(self):
        # rider1: 80.00 (default basket) + 100.00 = 180.00 held; rider2: 25.00 held
        self.a = self.deliver(self.rider_user, self.rider)
        self.b = self.deliver(self.rider_user, self.rider, quantity=4)
        self.c = self.deliver(self.rider2_user, self.rider2, quantity=1)
        self.staff_client = self.as_user(self.staff)

    def remit(self, **body):
        return self.staff_client.post(REMIT, body, format="json")

    def held_ids(self, rider):
        return set(
            Payment.objects.filter(verified_by=rider.user, remittance__isnull=True, payment_status="paid")
            .values_list("order_id", flat=True)
        )


class HeldCashTests(CashBase):
    def test_amount_held_per_rider(self):
        response = self.staff_client.get(HELD)
        self.assertEqual(response.status_code, 200)
        rows = {row["rider"]["id"]: row for row in response.json()}
        self.assertEqual(set(rows), {self.rider.pk, self.rider2.pk})

        first = rows[self.rider.pk]
        self.assertEqual(first["rider"], {"id": self.rider.pk, "full_name": "Rider1 Tester", "vehicle_type": ""})
        self.assertEqual(first["orders_count"], 2)
        self.assertEqual(first["amount_held"], "180.00")
        self.assertEqual({(o["order_id"], o["amount"]) for o in first["orders"]}, {(self.a, "80.00"), (self.b, "100.00")})
        self.assertTrue(all(o["collected_at"] for o in first["orders"]))
        self.assertEqual(first["oldest_collected_at"], first["orders"][0]["collected_at"])

        second = rows[self.rider2.pk]
        self.assertEqual((second["orders_count"], second["amount_held"]), (1, "25.00"))

    def test_largest_balance_first(self):
        self.assertEqual([r["rider"]["id"] for r in self.staff_client.get(HELD).json()], [self.rider.pk, self.rider2.pk])

    def test_only_paid_unremitted_cod_collected_by_the_rider_counts(self):
        self.deliver(self.rider_user, self.rider, method="gcash")  # online: no cash handled
        self.deliver(self.rider_user, self.rider, finish=None)  # still out for delivery
        self.deliver(self.rider2_user, self.rider2, finish="failed")  # failed: no cash
        rows = {r["rider"]["id"]: r for r in self.staff_client.get(HELD).json()}
        self.assertEqual(rows[self.rider.pk]["amount_held"], "180.00")
        self.assertEqual(rows[self.rider2.pk]["amount_held"], "25.00")

    def test_riders_with_nothing_are_left_out(self):
        self.remit(rider=self.rider2.pk)
        self.assertEqual([r["rider"]["id"] for r in self.staff_client.get(HELD).json()], [self.rider.pk])
        self.remit(rider=self.rider.pk)
        self.assertEqual(self.staff_client.get(HELD).json(), [])

    def test_rider_filter(self):
        rows = self.staff_client.get(f"{HELD}?rider={self.rider2.pk}").json()
        self.assertEqual([r["rider"]["id"] for r in rows], [self.rider2.pk])
        self.assertEqual(self.staff_client.get(f"{HELD}?rider=999999").json(), [])
        for bad in ("abc", "0", "-1", "1.5", "9" * 30):
            response = self.staff_client.get(f"{HELD}?rider={bad}")
            self.assertEqual(response.status_code, 400, bad)
            self.assertIn("rider", response.data["errors"], bad)

    def test_permissions(self):
        for user in (self.customer, self.rider_user):
            self.assertEqual(self.as_user(user).get(HELD).status_code, 403, user.role)
        self.assertEqual(APIClient().get(HELD).status_code, 401)
        self.assertEqual(self.as_user(self.admin).get(HELD).status_code, 200)


class RemitTests(CashBase):
    def test_remit_everything_a_rider_holds(self):
        response = self.remit(rider=self.rider.pk, remarks="End of day")
        self.assertEqual(response.status_code, 201, response.data)
        body = response.json()
        self.assertEqual(body["amount"], "180.00")
        self.assertEqual(body["remarks"], "End of day")
        self.assertEqual(body["rider"], {"id": self.rider.pk, "full_name": "Rider1 Tester"})
        self.assertEqual(body["received_by"], {"id": self.staff.pk, "full_name": "Staff1 Tester"})
        self.assertEqual([(o["order_id"], o["amount"]) for o in body["orders"]], sorted([(self.a, "80.00"), (self.b, "100.00")]))
        self.assertTrue(body["received_at"])
        self.assertEqual(self.held_ids(self.rider), set())
        self.assertEqual(self.held_ids(self.rider2), {self.c})  # the other rider is untouched
        self.assertEqual(CashRemittance.objects.count(), 1)

    def test_partial_remit_leaves_the_rest_held(self):
        response = self.remit(rider=self.rider.pk, orders=[self.b])
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.json()["amount"], "100.00")
        self.assertEqual(self.held_ids(self.rider), {self.a})
        row = self.staff_client.get(f"{HELD}?rider={self.rider.pk}").json()[0]
        self.assertEqual(row["amount_held"], "80.00")

        rest = self.remit(rider=self.rider.pk)  # now everything that is left
        self.assertEqual(rest.json()["amount"], "80.00")
        self.assertEqual(CashRemittance.objects.count(), 2)

    def test_amount_is_always_computed_by_the_server(self):
        response = self.remit(rider=self.rider.pk, amount="1.00", received_at="2000-01-01T00:00:00Z", received_by=self.admin.pk)
        self.assertEqual(response.status_code, 201)
        remittance = CashRemittance.objects.get()
        self.assertEqual(remittance.amount, Decimal("180.00"))
        self.assertEqual(remittance.received_by, self.staff)
        self.assertGreater(remittance.received_at.year, 2020)

    def test_admin_can_remit_too(self):
        response = self.as_user(self.admin).post(REMIT, {"rider": self.rider2.pk}, format="json")
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json()["received_by"]["id"], self.admin.pk)

    def test_the_same_order_cannot_be_remitted_twice(self):
        first = self.remit(rider=self.rider.pk, orders=[self.a])
        self.assertEqual(first.status_code, 201)
        again = self.remit(rider=self.rider.pk, orders=[self.a])
        self.assertEqual(again.status_code, 400)
        self.assertEqual(again.data["code"], "validation_error")
        self.assertIn(f"#{self.a}", again.data["errors"]["orders"][0])
        self.assertIn("already remitted", again.data["errors"]["orders"][0])
        self.assertEqual(CashRemittance.objects.count(), 1)

    def test_another_riders_order_is_rejected(self):
        response = self.remit(rider=self.rider.pk, orders=[self.c])
        self.assertEqual(response.status_code, 400)
        self.assertIn(f"#{self.c}", response.data["errors"]["orders"][0])
        self.assertEqual(self.held_ids(self.rider2), {self.c})
        self.assertEqual(CashRemittance.objects.count(), 0)

    def test_unknown_or_non_cash_orders_are_rejected(self):
        online = self.deliver(self.rider_user, self.rider, method="gcash")
        for bad in (999999, online):
            response = self.remit(rider=self.rider.pk, orders=[bad])
            self.assertEqual(response.status_code, 400, bad)
            self.assertIn(f"#{bad}", response.data["errors"]["orders"][0])

    def test_all_or_nothing(self):
        response = self.remit(rider=self.rider.pk, orders=[self.a, self.c, self.b])
        self.assertEqual(response.status_code, 400)
        self.assertEqual(CashRemittance.objects.count(), 0)
        self.assertEqual(self.held_ids(self.rider), {self.a, self.b})
        self.assertEqual(Notification.objects.filter(title="Cash handed in").count(), 0)

    def test_nothing_to_remit_is_a_400(self):
        self.remit(rider=self.rider2.pk)
        empty = self.remit(rider=self.rider2.pk)
        self.assertEqual(empty.status_code, 400)
        self.assertIn("not holding any cash", empty.data["detail"])
        self.assertEqual(self.remit(rider=self.rider2.pk, orders=[]).status_code, 400)
        self.assertEqual(CashRemittance.objects.count(), 1)

    def test_input_validation(self):
        for body in ({}, {"rider": 999999}, {"rider": "x"}, {"rider": self.rider.pk, "orders": "49"},
                     {"rider": self.rider.pk, "orders": [0]}, {"rider": self.rider.pk, "orders": [10**30]},
                     {"rider": self.rider.pk, "remarks": "x" * 501}):
            response = self.remit(**body)
            self.assertEqual(response.status_code, 400, body)
            self.assertEqual(response.data["code"], "validation_error", body)
        self.assertEqual(CashRemittance.objects.count(), 0)

    def test_duplicate_ids_in_one_call_are_counted_once(self):
        response = self.remit(rider=self.rider.pk, orders=[self.a, self.a])
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json()["amount"], "80.00")

    def test_rider_and_customer_cannot_remit(self):
        for user in (self.rider_user, self.rider2_user, self.customer):
            response = self.as_user(user).post(REMIT, {"rider": self.rider.pk}, format="json")
            self.assertEqual(response.status_code, 403, user.role)
        self.assertEqual(APIClient().post(REMIT, {"rider": self.rider.pk}, format="json").status_code, 401)
        self.assertEqual(CashRemittance.objects.count(), 0)

    def test_rider_is_notified(self):
        with self.captureOnCommitCallbacks(execute=True):
            self.remit(rider=self.rider.pk)
        note = Notification.objects.get(recipient=self.rider_user, title="Cash handed in")
        self.assertEqual(note.message, "Cash handed in: PHP 180.00 confirmed.")
        self.assertFalse(Notification.objects.filter(recipient=self.rider2_user, title="Cash handed in").exists())
        self.assertTrue(any("Cash handed in" in m.subject for m in mail.outbox))

    def test_payment_exposes_the_remittance_to_staff_and_admin_only(self):
        remittance = self.remit(rider=self.rider.pk, orders=[self.a]).json()
        for user in (self.staff, self.admin):
            payment = self.as_user(user).get(f"/api/orders/{self.a}/").json()["payment"]
            self.assertEqual(payment["remittance"], remittance["id"], user.role)
            self.assertEqual(payment["remitted_at"], remittance["received_at"], user.role)
            self.assertIsNone(self.as_user(user).get(f"/api/orders/{self.b}/").json()["payment"]["remittance"])
            self.assertIsNone(self.as_user(user).get(f"/api/orders/{self.b}/").json()["payment"]["remitted_at"])
            listed = self.as_user(user).get(f"/api/payments/?order={self.a}").json()["results"][0]
            self.assertEqual(listed["remittance"], remittance["id"])
        for user, url in ((self.customer, f"/api/orders/{self.a}/"), (self.rider_user, f"/api/deliveries/my/{self.a}/?scope=history")):
            payment = self.as_user(user).get(url).json()["payment"]
            self.assertNotIn("remittance", payment, user.role)
            self.assertNotIn("remitted_at", payment, user.role)
        customer_payments = self.as_user(self.customer).get("/api/payments/").json()["results"]
        self.assertTrue(all("remittance" not in p for p in customer_payments))


class RemittanceHistoryTests(CashBase):
    def setUp(self):
        super().setUp()
        self.first = self.remit(rider=self.rider.pk, orders=[self.a]).json()  # 80.00
        self.second = self.remit(rider=self.rider2.pk).json()  # 25.00
        self.third = self.remit(rider=self.rider.pk).json()  # 100.00

    def ids(self, query=""):
        response = self.staff_client.get(f"{HISTORY}?{query}")
        self.assertEqual(response.status_code, 200, response.data)
        return [r["id"] for r in response.json()["results"]]

    def test_list_shape_and_default_order(self):
        body = self.staff_client.get(HISTORY).json()
        self.assertEqual(set(body), {"count", "next", "previous", "results"})
        self.assertEqual(body["count"], 3)
        self.assertEqual([r["id"] for r in body["results"]], [self.third["id"], self.second["id"], self.first["id"]])
        item = body["results"][2]
        self.assertEqual(set(item), {"id", "rider", "amount", "orders", "received_by", "received_at", "remarks"})
        self.assertEqual(item["amount"], "80.00")
        self.assertEqual(item["orders"], [{"order_id": self.a, "amount": "80.00"}])
        self.assertEqual(item["rider"]["full_name"], "Rider1 Tester")
        self.assertEqual(item["received_by"]["full_name"], "Staff1 Tester")

    def test_filters_and_ordering(self):
        self.assertEqual(set(self.ids(f"rider={self.rider.pk}")), {self.first["id"], self.third["id"]})
        self.assertEqual(self.ids(f"rider={self.rider2.pk}"), [self.second["id"]])
        self.assertEqual(self.ids("ordering=amount"), [self.second["id"], self.first["id"], self.third["id"]])
        self.assertEqual(self.ids("ordering=-amount"), [self.third["id"], self.first["id"], self.second["id"]])
        self.assertEqual(len(self.ids("date_from=2000-01-01&date_to=2999-01-01")), 3)
        self.assertEqual(self.ids("date_to=2000-01-01"), [])
        self.assertEqual(self.ids("date_from=2999-01-01"), [])
        self.assertEqual(self.staff_client.get(f"{HISTORY}?rider=abc").status_code, 400)
        self.assertEqual(self.staff_client.get(f"{HISTORY}?date_from=nope").status_code, 400)
        self.assertEqual(len(self.staff_client.get(f"{HISTORY}?page_size=2").json()["results"]), 2)

    def test_detail(self):
        response = self.staff_client.get(f"{HISTORY}{self.third['id']}/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["amount"], "100.00")
        self.assertEqual(self.staff_client.get(f"{HISTORY}999999/").status_code, 404)

    def test_remittances_cannot_be_edited_or_deleted(self):
        url = f"{HISTORY}{self.first['id']}/"
        admin = self.as_user(self.admin)
        for method in (admin.put, admin.patch, admin.delete, admin.post):
            self.assertEqual(method(url, {"amount": "1.00"}, format="json").status_code, 405)
        self.assertEqual(admin.post(HISTORY, {"amount": "1.00"}, format="json").status_code, 405)
        self.assertEqual(CashRemittance.objects.get(pk=self.first["id"]).amount, Decimal("80.00"))

    def test_staff_and_admin_only(self):
        for user in (self.customer, self.rider_user):
            self.assertEqual(self.as_user(user).get(HISTORY).status_code, 403, user.role)
            self.assertEqual(self.as_user(user).get(f"{HISTORY}{self.first['id']}/").status_code, 403)
        self.assertEqual(APIClient().get(HISTORY).status_code, 401)
        self.assertEqual(self.as_user(self.admin).get(HISTORY).status_code, 200)


class RiderOwnBalanceTests(CashBase):
    def test_rider_sees_only_their_own_cash(self):
        body = self.as_user(self.rider_user).get(MINE).json()
        self.assertEqual(body["amount_held"], "180.00")
        self.assertEqual({(o["order_id"], o["amount"]) for o in body["orders"]}, {(self.a, "80.00"), (self.b, "100.00")})
        self.assertTrue(all(o["collected_at"] for o in body["orders"]))
        self.assertEqual(body["recent_remittances"], [])
        other = self.as_user(self.rider2_user).get(MINE).json()
        self.assertEqual(other["amount_held"], "25.00")
        self.assertEqual([o["order_id"] for o in other["orders"]], [self.c])

    def test_balance_and_recent_hand_ins_after_a_remittance(self):
        self.remit(rider=self.rider.pk, orders=[self.a])
        done = self.remit(rider=self.rider.pk).json()
        body = self.as_user(self.rider_user).get(MINE).json()
        self.assertEqual(body["amount_held"], "0.00")
        self.assertEqual(body["orders"], [])
        recent = body["recent_remittances"]
        self.assertEqual([r["id"] for r in recent][0], done["id"])  # newest first
        self.assertEqual(len(recent), 2)
        self.assertEqual(set(recent[0]), {"id", "amount", "received_at", "received_by"})
        self.assertEqual(recent[0]["amount"], "100.00")
        self.assertEqual(recent[0]["received_by"], "Staff1 Tester")
        # another rider's hand-ins never show up here
        self.assertEqual(self.as_user(self.rider2_user).get(MINE).json()["recent_remittances"], [])

    def test_rider_with_no_cash(self):
        self.remit(rider=self.rider2.pk)
        body = self.as_user(self.rider2_user).get(MINE).json()
        self.assertEqual((body["amount_held"], body["orders"]), ("0.00", []))

    def test_only_riders_can_read_it(self):
        for user in (self.customer, self.staff, self.admin):
            self.assertEqual(self.as_user(user).get(MINE).status_code, 403, user.role)
        self.assertEqual(APIClient().get(MINE).status_code, 401)

    def test_rider_cannot_remit_or_see_the_staff_screens(self):
        client = self.as_user(self.rider_user)
        self.assertEqual(client.post(REMIT, {"rider": self.rider.pk}, format="json").status_code, 403)
        self.assertEqual(client.get(HELD).status_code, 403)


class DashboardCashTests(CashBase):
    def test_cash_outstanding_follows_remittances(self):
        self.assertEqual(self.staff_client.get("/api/orders/dashboard/").json()["cash_outstanding"], "205.00")
        self.remit(rider=self.rider.pk, orders=[self.a])
        self.assertEqual(self.staff_client.get("/api/orders/dashboard/").json()["cash_outstanding"], "125.00")
        self.remit(rider=self.rider.pk)
        self.remit(rider=self.rider2.pk)
        self.assertEqual(self.as_user(self.admin).get("/api/orders/dashboard/").json()["cash_outstanding"], "0.00")

    def test_cash_outstanding_is_zero_with_no_orders(self):
        Order.objects.all().delete()
        self.assertEqual(self.staff_client.get("/api/orders/dashboard/").json()["cash_outstanding"], "0.00")


class CashReportTests(CashBase):
    URL = "/api/reports/cash/"

    def test_per_rider_collected_remitted_outstanding(self):
        self.remit(rider=self.rider.pk, orders=[self.a])  # hands in 80.00 of 180.00
        data = self.as_user(self.admin).get(self.URL).json()
        self.assertEqual(data["report"], "cash")
        self.assertEqual([c["key"] for c in data["columns"]], ["rider", "collected", "remitted", "outstanding"])
        rows = {r["rider_id"]: r for r in data["rows"]}
        one = rows[self.rider.pk]
        self.assertEqual((one["collected"], one["remitted"], one["outstanding"]), ("180.00", "80.00", "100.00"))
        two = rows[self.rider2.pk]
        self.assertEqual((two["collected"], two["remitted"], two["outstanding"]), ("25.00", "0.00", "25.00"))
        self.assertEqual(data["totals"], {"collected": "205.00", "remitted": "80.00", "outstanding": "125.00"})

    def test_date_range_limits_collected_and_remitted_but_not_outstanding(self):
        self.remit(rider=self.rider.pk, orders=[self.a])
        past = self.as_user(self.admin).get(self.URL, {"date_from": "2000-01-01", "date_to": "2000-01-31"}).json()
        self.assertEqual(past["totals"], {"collected": "0.00", "remitted": "0.00", "outstanding": "125.00"})

    def test_admin_only_and_exports(self):
        for user in (self.staff, self.customer, self.rider_user):
            self.assertEqual(self.as_user(user).get(self.URL).status_code, 403, user.role)
        admin = self.as_user(self.admin)
        pdf = admin.get(self.URL, {"format": "pdf"})
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        xlsx = admin.get(self.URL, {"format": "xlsx"})
        sheet = load_workbook(io.BytesIO(xlsx.content)).active
        values = [[c.value for c in row] for row in sheet.iter_rows()]
        self.assertTrue(any(v[0] == "Total" for v in values))
        self.assertIn("cash_", xlsx["Content-Disposition"])
